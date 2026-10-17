"""Text and embedded-image extraction from uploaded PDF / DOCX / TXT / MD files.

Edit when: changing how uploaded reports are read.
"""

import io
from typing import List, Dict

from docx import Document
from pypdf import PdfReader


def extract_text_from_uploaded_file(uploaded_file) -> str:
    """Extract plain text from uploaded PDF, DOCX, TXT, or MD files."""
    if uploaded_file is None:
        return ""

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join([page.extract_text() or "" for page in reader.pages])

    elif filename.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

    elif filename.endswith((".txt", ".md")):
        return file_bytes.decode("utf-8", errors="ignore")

    return ""


def extract_images_from_report(uploaded_file) -> List[Dict[str, bytes]]:
    """
    Extract embedded images from an uploaded PDF or DOCX report so users can reuse
    original diagrams, charts, or photos directly inside their slides.
    """
    extracted: List[Dict[str, bytes]] = []
    if uploaded_file is None:
        return extracted

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    try:
        if filename.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(file_bytes))
            img_counter = 1
            for page in reader.pages:
                for image_file_object in page.images:
                    extracted.append({
                        "name": f"Report_Image_{img_counter} ({image_file_object.name})",
                        "bytes": image_file_object.data
                    })
                    img_counter += 1

        elif filename.endswith(".docx"):
            doc = Document(io.BytesIO(file_bytes))
            img_counter = 1
            for rel in doc.part.rels.values():
                if "image" in rel.target_ref:
                    extracted.append({
                        "name": f"Report_Image_{img_counter}",
                        "bytes": rel.target_part.blob
                    })
                    img_counter += 1
    except Exception:
        pass

    return extracted

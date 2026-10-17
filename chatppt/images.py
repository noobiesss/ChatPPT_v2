"""Slide image engine: Unsplash search, AI image fallback, and fitted picture insertion.

Edit when: changing image sources, query logic, or how pictures are placed on a slide.
"""

import io
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from typing import List, Dict, Optional

from google import genai
from google.genai import types
from PIL import Image
from pptx.util import Inches

from chatppt.schemas import IMAGE_MODE_PLACEHOLDER, IMAGE_MODE_UNSPLASH_FIRST


def _derive_fallback_keywords(text: str) -> str:
    """Extract 2-3 clean ASCII words from a title for Unsplash search fallback, or return a safe default."""
    words = [w.lower() for w in re.findall(r"[A-Za-z]{4,}", text or "") if w.lower() not in {
        "that", "with", "from", "this", "have", "were", "their", "across", "through", "into", "over", "more", "than"
    }]
    if len(words) >= 2:
        return " ".join(words[:3])
    return "business technology office"


def _clean_unsplash_key(raw_key: Optional[str]) -> str:
    """Normalize an Unsplash Access Key and strip accidental 'Client-ID ' or 'Bearer ' prefixes."""
    key = (raw_key or "").strip()
    for prefix in ("Client-ID ", "client-id ", "Bearer ", "bearer "):
        if key.startswith(prefix):
            key = key[len(prefix):].strip()
    return key


def _query_unsplash_once(query_text: str, access_key: str) -> Optional[bytes]:
    """Run a single Unsplash search query and download the first landscape photo."""
    q = urllib.parse.quote(query_text.strip())
    if not q:
        return None
    api_url = (
        f"https://api.unsplash.com/search/photos"
        f"?query={q}&per_page=1&orientation=landscape&content_filter=high"
    )
    req = urllib.request.Request(
        api_url,
        headers={
            "Authorization": f"Client-ID {access_key}",
            "Accept-Version": "v1",
            "User-Agent": "ChatPPT-V2",
        },
    )
    with urllib.request.urlopen(req, timeout=10) as resp:
        data = json.loads(resp.read().decode("utf-8"))
        results = data.get("results", [])
        if not results:
            return None
        img_url = results[0]["urls"]["regular"]

    img_req = urllib.request.Request(img_url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(img_req, timeout=12) as img_resp:
        return img_resp.read()


def fetch_unsplash_image_bytes(search_keywords: str, unsplash_access_key: Optional[str]) -> Optional[bytes]:
    """
    Query Unsplash Search API for a high-resolution landscape stock photo matching search_keywords.
    If a multi-word query yields 0 results, retries with the first 2 keywords before returning None.
    """
    access_key = _clean_unsplash_key(unsplash_access_key)
    if not access_key or not (search_keywords or "").strip():
        return None

    cleaned_kw = re.sub(r"[^\w\s-]", " ", search_keywords).strip()
    words = [w for w in cleaned_kw.split() if w]
    if not words:
        return None

    candidate_queries = [" ".join(words)]
    if len(words) > 2:
        short_q = " ".join(words[:2])
        if short_q not in candidate_queries:
            candidate_queries.append(short_q)

    for q in candidate_queries:
        try:
            img_bytes = _query_unsplash_once(q, access_key)
            if img_bytes:
                return img_bytes
        except urllib.error.HTTPError as http_err:
            # Stop retrying on auth/rate-limit errors (401/403/429)
            if http_err.code in (401, 403, 429):
                return None
        except Exception:
            continue

    return None


def generate_ai_image_bytes(
    visual_prompt: str,
    gemini_api_key: Optional[str] = None,
    use_imagen: bool = False
) -> Optional[bytes]:
    """
    Generate a 4:3 conceptual illustration for a slide when Unsplash has no photo or no key.
    Tries Google Imagen 3 first (if enabled and a Gemini API key is available),
    and falls back to Pollinations AI stream.
    """
    enhanced_prompt = (
        f"{visual_prompt}, clean modern corporate presentation illustration, "
        "minimalist high-contrast aesthetic, no text, no watermarks"
    )

    if use_imagen and gemini_api_key:
        try:
            client = genai.Client(api_key=gemini_api_key)
            result = client.models.generate_images(
                model="imagen-3.0-generate-002",
                prompt=enhanced_prompt,
                config=types.GenerateImagesConfig(
                    number_of_images=1,
                    output_mime_type="image/jpeg",
                    aspect_ratio="4:3",
                ),
            )
            for generated_image in result.generated_images:
                return generated_image.image.image_bytes
        except Exception:
            pass

    try:
        encoded_prompt = urllib.parse.quote(enhanced_prompt[:300])
        url = f"https://image.pollinations.ai/prompt/{encoded_prompt}?width=800&height=600&nologo=true&seed=42"
        req = urllib.request.Request(
            url,
            headers={"User-Agent": "Mozilla/5.0", "Accept": "image/*,*/*;q=0.8"},
        )
        with urllib.request.urlopen(req, timeout=18) as response:
            return response.read()
    except Exception:
        return None


def resolve_slide_image_bytes(
    slide_data: dict,
    gemini_api_key: Optional[str],
    use_imagen: bool,
    unsplash_access_key: str,
    extracted_report_images: Optional[List[Dict[str, bytes]]] = None,
    default_image_mode: str = IMAGE_MODE_UNSPLASH_FIRST,
) -> Optional[bytes]:
    """
    Two-Option Image Resolver:
    1. IMAGE_MODE_UNSPLASH_FIRST (default): Query Unsplash stock photos first; if no key or no match,
       automatically fall back to AI image generation.
    2. IMAGE_MODE_PLACEHOLDER: Return None so the slide renders a clean reserved photo placeholder box.
    """
    src_mode = slide_data.get("image_source_mode") or default_image_mode
    if src_mode == IMAGE_MODE_PLACEHOLDER:
        return None

    image_spec = slide_data.get("image") or {}
    action_title = slide_data.get("action_title", "Corporate illustration")
    visual_prompt = (image_spec.get("visual_prompt") or "").strip() or action_title
    search_keywords = (image_spec.get("search_keywords") or "").strip() or _derive_fallback_keywords(action_title)

    # Option 1 (default): Prioritize Unsplash photo; if unavailable or no result, fall back to AI generation
    img_bytes = fetch_unsplash_image_bytes(search_keywords, unsplash_access_key)
    if img_bytes:
        return img_bytes

    return generate_ai_image_bytes(visual_prompt, gemini_api_key=gemini_api_key, use_imagen=use_imagen)


def add_picture_fitted(
    slide,
    img_bytes: bytes,
    box_left: float,
    box_top: float,
    box_width: float,
    box_height: float
) -> bool:
    """
    Insert an image into a bounding box (in Inches) while preserving its native aspect ratio
    and centering it inside the box. Converts unsupported modes (WEBP/RGBA) to RGB automatically.
    """
    try:
        with Image.open(io.BytesIO(img_bytes)) as pil_img:
            if pil_img.mode in ("RGBA", "P") or pil_img.format not in ("JPEG", "PNG"):
                rgb_img = pil_img.convert("RGB")
                buf = io.BytesIO()
                rgb_img.save(buf, format="JPEG", quality=92)
                img_bytes = buf.getvalue()
                img_w, img_h = rgb_img.size
            else:
                img_w, img_h = pil_img.size

        img_ratio = img_w / float(img_h)
        box_ratio = box_width / float(box_height)

        if img_ratio > box_ratio:
            final_w = box_width
            final_h = box_width / img_ratio
        else:
            final_h = box_height
            final_w = box_height * img_ratio

        offset_left = box_left + (box_width - final_w) / 2.0
        offset_top = box_top + (box_height - final_h) / 2.0

        slide.shapes.add_picture(
            io.BytesIO(img_bytes),
            Inches(offset_left),
            Inches(offset_top),
            width=Inches(final_w),
            height=Inches(final_h),
        )
        return True
    except Exception:
        return False

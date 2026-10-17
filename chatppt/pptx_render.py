"""Slide drawing helpers and the 8 layout renderers (cover, divider, cards, columns, KPI, chart, image split).

Edit when: changing how a given slide layout looks. Does not touch template inheritance or the UI.
"""

from typing import Dict, Optional

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from chatppt.images import add_picture_fitted
from chatppt.pptx_template import _ensure_white_text_contrast
from chatppt.schemas import ChartType


def apply_slide_fonts(slide, theme: Dict[str, object]) -> None:
    """Apply the template's heading and body fonts across all text frames on a rendered slide."""
    font_heading = theme.get("font_heading")
    font_body = theme.get("font_body") or font_heading
    if not font_heading and not font_body:
        return

    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for para in shape.text_frame.paragraphs:
            is_heading = bool(para.font.bold) or (para.font.size is not None and para.font.size >= Pt(16))
            target_font = font_heading if is_heading else font_body
            if target_font:
                para.font.name = target_font
                for run in para.runs:
                    run.font.name = target_font


def set_solid_background(slide, color: RGBColor) -> None:
    """Apply a solid background color to a slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def clear_existing_slides(prs: Presentation) -> None:
    """Remove sample slides from an uploaded .pptx template while keeping Slide Masters intact."""
    for i in range(len(prs.slides) - 1, -1, -1):
        r_id = prs.slides._sldIdLst[i].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[i]


def draw_card_shape(slide, left: float, top: float, width: float, height: float, theme: Dict[str, object]):
    """Draw a styled rounded rectangle card container."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.adjustments[0] = 0.04
    card.fill.solid()
    card.fill.fore_color.rgb = theme["card_bg"]
    card.line.color.rgb = theme["card_border"]
    card.line.width = Pt(1)
    return card


def draw_header_and_takeaway(
    slide,
    slide_data: dict,
    slide_idx: int,
    total_slides: int,
    theme: Dict[str, object],
    apply_bg: bool = True
) -> None:
    """Render the top accent bar, Action Title, bottom takeaway banner, and slide number."""
    if apply_bg:
        set_solid_background(slide, theme["bg_content"])
        # Top accent strip (only drawn on built-in themes so custom template top banners aren't covered)
        top_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12)
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = theme["accent"]
        top_bar.line.fill.background()

    # Action Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.7), Inches(1.05))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = slide_data.get("action_title", "")
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = theme["title_content"]

    # Accent underline below title
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.52), Inches(11.7), Inches(0.03)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = theme["accent"]
    divider.line.fill.background()

    # Bottom Takeaway Banner ("So What?")
    takeaway = (slide_data.get("bottom_takeaway") or "").strip()
    if takeaway:
        banner = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(10.8), Inches(0.68)
        )
        banner.adjustments[0] = 0.15
        banner.fill.solid()
        banner.fill.fore_color.rgb = theme["takeaway_bg"]
        banner.line.fill.background()

        tb_take = slide.shapes.add_textbox(Inches(1.0), Inches(6.42), Inches(10.4), Inches(0.55))
        tf_take = tb_take.text_frame
        tf_take.word_wrap = True
        p_take = tf_take.paragraphs[0]
        p_take.text = f"💡 Key Takeaway: {takeaway}"
        p_take.font.size = Pt(13.5)
        p_take.font.bold = True
        p_take.font.color.rgb = theme["takeaway_text"]

    # Slide Number Footer
    footer_box = slide.shapes.add_textbox(Inches(11.7), Inches(6.45), Inches(0.8), Inches(0.45))
    tf_footer = footer_box.text_frame
    p_footer = tf_footer.paragraphs[0]
    p_footer.text = f"{slide_idx + 1} / {total_slides}"
    p_footer.alignment = PP_ALIGN.RIGHT
    p_footer.font.size = Pt(11)
    p_footer.font.color.rgb = theme["accent"]


def render_cover_slide(slide, slide_data: dict, theme: Dict[str, object], apply_bg: bool = True) -> None:
    """Render Layout 1: Cover Slide."""
    if apply_bg:
        set_solid_background(slide, theme["bg_cover"])

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(2.2), Inches(1.4), Inches(0.14)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = theme["accent"]
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.5), Inches(11.0), Inches(2.2))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = slide_data.get("action_title", "Presentation Title")
    p_title.font.size = Pt(42)
    p_title.font.bold = True
    p_title.font.color.rgb = theme["title_cover"]

    subtitle = (slide_data.get("subtitle_or_context") or "").strip()
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(1.1), Inches(4.8), Inches(10.5), Inches(1.4))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.size = Pt(20)
        p_sub.font.color.rgb = theme["subtitle_cover"]


def render_section_divider_slide(slide, slide_data: dict, theme: Dict[str, object], apply_bg: bool = True) -> None:
    """Render Layout 2: Section Divider Slide."""
    if apply_bg:
        set_solid_background(slide, theme["bg_section"])

    accent_pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(2.6), Inches(0.9), Inches(0.12)
    )
    accent_pill.fill.solid()
    accent_pill.fill.fore_color.rgb = theme["accent"]
    accent_pill.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.9), Inches(11.0), Inches(1.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = slide_data.get("action_title", "Section Title")
    p_title.font.size = Pt(38)
    p_title.font.bold = True
    p_title.font.color.rgb = theme["title_cover"]

    context = (slide_data.get("subtitle_or_context") or "").strip()
    if context:
        sub_box = slide.shapes.add_textbox(Inches(1.1), Inches(4.5), Inches(10.0), Inches(1.2))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = context
        p_sub.font.size = Pt(18)
        p_sub.font.color.rgb = theme["subtitle_cover"]


def render_bullet_cards_slide(slide, slide_data: dict, theme: Dict[str, object]) -> None:
    """Render Layout 3: Stacked Horizontal Structured Cards."""
    points = (slide_data.get("points") or [])[:4]
    count = max(len(points), 1)
    avail_h = 4.35
    gap = 0.16
    card_h = (avail_h - gap * (count - 1)) / count

    for idx, pt in enumerate(points):
        top = 1.78 + idx * (card_h + gap)
        draw_card_shape(slide, 0.8, top, 11.7, card_h, theme)

        # Left number badge
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1.05), Inches(top + (card_h - 0.46) / 2), Inches(0.46), Inches(0.46)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = _ensure_white_text_contrast(theme["accent"])
        badge.line.fill.background()
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = str(idx + 1)
        p_b.alignment = PP_ALIGN.CENTER
        p_b.font.size = Pt(14)
        p_b.font.bold = True
        p_b.font.color.rgb = RGBColor(255, 255, 255)

        # Heading & Detail text
        tb = slide.shapes.add_textbox(Inches(1.7), Inches(top + 0.08), Inches(10.5), Inches(card_h - 0.16))
        tf = tb.text_frame
        tf.word_wrap = True
        p_head = tf.paragraphs[0]
        p_head.text = pt.get("heading", "")
        p_head.font.size = Pt(17)
        p_head.font.bold = True
        p_head.font.color.rgb = theme["title_content"]
        p_head.space_after = Pt(3)

        p_det = tf.add_paragraph()
        p_det.text = pt.get("detail", "")
        p_det.font.size = Pt(14)
        p_det.font.color.rgb = theme["body_content"]


def render_three_cards_slide(slide, slide_data: dict, theme: Dict[str, object]) -> None:
    """Render Layout 4: 3 Parallel Vertical Pillar Cards."""
    points = (slide_data.get("points") or [])[:3]
    while len(points) < 3:
        points.append({"heading": f"Pillar {len(points) + 1}", "detail": "Key supporting detail."})

    card_w = 3.7
    gap = 0.3
    card_h = 4.3

    for idx, pt in enumerate(points):
        left = 0.8 + idx * (card_w + gap)
        top = 1.8
        draw_card_shape(slide, left, top, card_w, card_h, theme)

        # Accent header bar inside each pillar card
        strip = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left + 0.25), Inches(top + 0.28), Inches(0.8), Inches(0.08)
        )
        strip.fill.solid()
        strip.fill.fore_color.rgb = theme["accent"]
        strip.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.5), Inches(card_w - 0.5), Inches(card_h - 0.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p_head = tf.paragraphs[0]
        p_head.text = pt.get("heading", "")
        p_head.font.size = Pt(19)
        p_head.font.bold = True
        p_head.font.color.rgb = theme["title_content"]
        p_head.space_after = Pt(12)

        p_det = tf.add_paragraph()
        p_det.text = pt.get("detail", "")
        p_det.font.size = Pt(14.5)
        p_det.font.color.rgb = theme["body_content"]


def render_two_column_slide(slide, slide_data: dict, theme: Dict[str, object]) -> None:
    """Render Layout 5: Side-by-Side Two-Column Comparison."""
    left_col = slide_data.get("left_column") or {"column_title": "Column A", "items": []}
    right_col = slide_data.get("right_column") or {"column_title": "Column B", "items": []}

    col_w = 5.65
    gap = 0.4
    card_h = 4.3

    for idx, col_data in enumerate([left_col, right_col]):
        left = 0.8 + idx * (col_w + gap)
        top = 1.8
        draw_card_shape(slide, left, top, col_w, card_h, theme)

        # Column Header Banner
        header_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.2), Inches(top + 0.2), Inches(col_w - 0.4), Inches(0.55)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = _ensure_white_text_contrast(theme["accent"]) if idx == 1 else theme["takeaway_bg"]
        header_shape.line.fill.background()

        tf_h = header_shape.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = col_data.get("column_title", "")
        p_h.alignment = PP_ALIGN.CENTER
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = RGBColor(255, 255, 255)

        # Column Items
        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.9), Inches(col_w - 0.5), Inches(card_h - 1.1))
        tf = tb.text_frame
        tf.word_wrap = True

        for item_idx, item in enumerate((col_data.get("items") or [])[:4]):
            p_head = tf.paragraphs[0] if item_idx == 0 else tf.add_paragraph()
            p_head.text = f"▪ {item.get('heading', '')}"
            p_head.font.size = Pt(15.5)
            p_head.font.bold = True
            p_head.font.color.rgb = theme["title_content"]
            p_head.space_after = Pt(2)

            p_det = tf.add_paragraph()
            p_det.text = f"   {item.get('detail', '')}"
            p_det.font.size = Pt(13.5)
            p_det.font.color.rgb = theme["body_content"]
            p_det.space_after = Pt(10)


def render_kpi_metrics_slide(slide, slide_data: dict, theme: Dict[str, object]) -> None:
    """Render Layout 6: Hero KPI Metric Cards."""
    kpis = (slide_data.get("kpis") or [])[:4]
    count = max(len(kpis), 1)
    total_w = 11.7
    gap = 0.28
    card_w = (total_w - gap * (count - 1)) / count
    card_h = 4.2

    for idx, kpi in enumerate(kpis):
        left = 0.8 + idx * (card_w + gap)
        top = 1.85
        draw_card_shape(slide, left, top, card_w, card_h, theme)

        tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.5), Inches(card_w - 0.4), Inches(card_h - 0.8))
        tf = tb.text_frame
        tf.word_wrap = True

        # Big Hero Metric Number
        p_val = tf.paragraphs[0]
        p_val.text = kpi.get("value", "")
        p_val.alignment = PP_ALIGN.CENTER
        p_val.font.size = Pt(46 if count <= 3 else 38)
        p_val.font.bold = True
        p_val.font.color.rgb = theme["accent"]
        p_val.space_after = Pt(12)

        # Metric Label
        p_lbl = tf.add_paragraph()
        p_lbl.text = kpi.get("label", "")
        p_lbl.alignment = PP_ALIGN.CENTER
        p_lbl.font.size = Pt(18)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = theme["title_content"]
        p_lbl.space_after = Pt(10)

        # Metric Context
        p_ctx = tf.add_paragraph()
        p_ctx.text = kpi.get("context", "")
        p_ctx.alignment = PP_ALIGN.CENTER
        p_ctx.font.size = Pt(13.5)
        p_ctx.font.color.rgb = theme["body_content"]


def render_chart_slide(slide, slide_data: dict, theme: Dict[str, object]) -> None:
    """Render Layout 7: Native Editable PowerPoint Chart (Left) + Key Insights (Right)."""
    chart_spec = slide_data.get("chart") or {}
    categories = chart_spec.get("categories") or ["A", "B", "C"]
    values = chart_spec.get("values") or [10.0, 20.0, 30.0]
    min_len = min(len(categories), len(values))
    categories = categories[:min_len]
    values = values[:min_len]

    # Left Card Frame for Chart
    draw_card_shape(slide, 0.8, 1.8, 6.4, 4.3, theme)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(chart_spec.get("series_name", "Series 1"), values)

    c_type_str = chart_spec.get("chart_type", "bar")
    if c_type_str == ChartType.LINE.value:
        xl_type = XL_CHART_TYPE.LINE_MARKERS
    elif c_type_str == ChartType.PIE.value:
        xl_type = XL_CHART_TYPE.PIE
    else:
        xl_type = XL_CHART_TYPE.COLUMN_CLUSTERED

    chart_frame = slide.shapes.add_chart(
        xl_type, Inches(1.0), Inches(1.95), Inches(6.0), Inches(4.0), chart_data
    )
    chart = chart_frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False

    if chart_spec.get("chart_title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = chart_spec["chart_title"]

    # Right Side Commentary Cards
    points = (slide_data.get("points") or [])[:3]
    count = max(len(points), 1)
    card_h = (4.3 - 0.15 * (count - 1)) / count

    for idx, pt in enumerate(points):
        top = 1.8 + idx * (card_h + 0.15)
        draw_card_shape(slide, 7.45, top, 5.05, card_h, theme)

        tb = slide.shapes.add_textbox(Inches(7.65), Inches(top + 0.1), Inches(4.65), Inches(card_h - 0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_head = tf.paragraphs[0]
        p_head.text = pt.get("heading", "")
        p_head.font.size = Pt(16)
        p_head.font.bold = True
        p_head.font.color.rgb = theme["title_content"]
        p_head.space_after = Pt(4)

        p_det = tf.add_paragraph()
        p_det.text = pt.get("detail", "")
        p_det.font.size = Pt(13)
        p_det.font.color.rgb = theme["body_content"]


def render_image_split_slide(
    slide,
    slide_data: dict,
    theme: Dict[str, object],
    img_bytes: Optional[bytes] = None
) -> None:
    """Render Layout 8: Visual Image + Caption on one side, Structured Points on the other."""
    image_info = slide_data.get("image") or {}
    placement = image_info.get("placement", "left")
    caption_text = image_info.get("caption", "")
    visual_prompt = image_info.get("visual_prompt", "Visual Illustration")
    search_kw = image_info.get("search_keywords", "")

    if placement == "right":
        text_left, img_left = 0.8, 6.95
    else:
        img_left, text_left = 0.8, 6.55

    img_w, img_h = 5.55, 3.65
    img_top = 1.8

    # Framed Container for Image + Caption
    draw_card_shape(slide, img_left, img_top, img_w, 4.3, theme)

    inserted = False
    if img_bytes:
        inserted = add_picture_fitted(
            slide,
            img_bytes,
            box_left=img_left + 0.15,
            box_top=img_top + 0.15,
            box_width=img_w - 0.3,
            box_height=img_h - 0.15,
        )

    if not inserted:
        # Styled Reserved Photo Placeholder
        ph_box = slide.shapes.add_textbox(
            Inches(img_left + 0.3), Inches(img_top + 0.6), Inches(img_w - 0.6), Inches(img_h - 0.7)
        )
        tf_ph = ph_box.text_frame
        tf_ph.word_wrap = True
        p_ph = tf_ph.paragraphs[0]
        p_ph.text = f"🖼️ [ Reserved Photo Placeholder ]\n\nKeywords: {search_kw}\nPrompt: {visual_prompt}"
        p_ph.alignment = PP_ALIGN.CENTER
        p_ph.font.size = Pt(12.5)
        p_ph.font.color.rgb = theme["body_content"]

    if caption_text:
        cap_box = slide.shapes.add_textbox(
            Inches(img_left + 0.2), Inches(img_top + img_h + 0.05), Inches(img_w - 0.4), Inches(0.5)
        )
        tf_cap = cap_box.text_frame
        tf_cap.word_wrap = True
        p_cap = tf_cap.paragraphs[0]
        p_cap.text = f"▲ {caption_text}"
        p_cap.alignment = PP_ALIGN.CENTER
        p_cap.font.size = Pt(11.5)
        p_cap.font.italic = True
        p_cap.font.color.rgb = theme["accent"]

    # Opposite Side: Structured Key Takeaways
    points = (slide_data.get("points") or [])[:3]
    count = max(len(points), 1)
    card_h = (4.3 - 0.15 * (count - 1)) / count

    for idx, pt in enumerate(points):
        top = 1.8 + idx * (card_h + 0.15)
        draw_card_shape(slide, text_left, top, 5.95, card_h, theme)

        tb = slide.shapes.add_textbox(Inches(text_left + 0.22), Inches(top + 0.12), Inches(5.5), Inches(card_h - 0.24))
        tf = tb.text_frame
        tf.word_wrap = True

        p_head = tf.paragraphs[0]
        p_head.text = pt.get("heading", "")
        p_head.font.size = Pt(16.5)
        p_head.font.bold = True
        p_head.font.color.rgb = theme["title_content"]
        p_head.space_after = Pt(4)

        p_det = tf.add_paragraph()
        p_det.text = pt.get("detail", "")
        p_det.font.size = Pt(13.5)
        p_det.font.color.rgb = theme["body_content"]

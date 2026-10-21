"""Offline smoke test for ChatPPT V2 (no API keys, no network).

Run from the project root after ANY code change:   python smoke_test.py
It checks that (1) every module imports, (2) a deck with all 11 layouts builds and re-opens,
(3) custom .pptx template inheritance still strips border frames and keeps template fonts,
(4) a few pure helpers still behave. Exit code 0 = all good.
"""

import importlib
import io
import pkgutil
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def check_imports() -> None:
    import chatppt

    # Import every submodule recursively, so a syntax error or a broken import anywhere fails fast
    names = [m.name for m in pkgutil.walk_packages(chatppt.__path__, "chatppt.")]
    for name in names:
        importlib.import_module(name)
    importlib.import_module("app")
    from PIL import Image

    from chatppt.ui.common import LOGO_ICON_PATH, LOGO_PATH

    with Image.open(LOGO_PATH) as logo, Image.open(LOGO_ICON_PATH) as icon:
        assert logo.mode == "RGBA" and icon.size[0] == icon.size[1], "logo assets missing or malformed"
    print(f"[ok] imported {len(names)} modules + app; logo assets present")


def sample_slides() -> list:
    # One slide per layout type (11 in total), shaped like a sanitized outline; other checks index into it
    pts = [{"heading": f"Point {i}", "detail": "Supporting detail text for this point."} for i in range(1, 4)]
    return [
        {"slide_number": 1, "layout_type": "cover", "action_title": "Deck Title", "subtitle_or_context": "Subtitle"},
        {"slide_number": 2, "layout_type": "section_divider", "action_title": "Section", "subtitle_or_context": "Overview"},
        {"slide_number": 3, "layout_type": "bullet_cards", "action_title": "Bullet cards title", "points": pts, "bottom_takeaway": "So what."},
        {"slide_number": 4, "layout_type": "three_cards", "action_title": "Three cards title", "points": pts, "bottom_takeaway": "So what."},
        {"slide_number": 5, "layout_type": "two_column", "action_title": "Two columns title",
         "left_column": {"column_title": "Before", "items": pts[:2]}, "right_column": {"column_title": "After", "items": pts[:2]},
         "bottom_takeaway": "So what."},
        {"slide_number": 6, "layout_type": "kpi_metrics", "action_title": "KPI title",
         "kpis": [{"value": "+62%", "label": "Growth", "context": "Year over year"}, {"value": "3.5x", "label": "Speed", "context": "Faster queries"}],
         "bottom_takeaway": "So what."},
        {"slide_number": 7, "layout_type": "chart_slide", "action_title": "Chart title", "points": pts[:2],
         "chart": {"chart_type": "bar", "chart_title": "Chart", "categories": ["A", "B", "C"], "series_name": "Rev ($M)", "values": [1.0, 2.0, 3.0]},
         "bottom_takeaway": "So what."},
        {"slide_number": 8, "layout_type": "image_split", "action_title": "Image split title", "points": pts[:2],
         "image": {"preferred_style": "stock_photo", "visual_prompt": "x", "search_keywords": "office", "caption": "Caption", "placement": "left"},
         "image_source_mode": "Reserve Photo Placeholder", "bottom_takeaway": "So what."},
        {"slide_number": 9, "layout_type": "timeline_flow", "action_title": "Timeline title",
         "timeline_steps": [{"stage_label": f"Q{i}", "title": f"Step {i}", "description": "What happens."} for i in range(1, 5)],
         "bottom_takeaway": "So what."},
        {"slide_number": 10, "layout_type": "comparison_table", "action_title": "Table title",
         "table_headers": ["Criterion", "Option A", "Option B"],
         "table_rows": [["Cost", "Low", "High"], ["Speed", "Fast", "Slow"], ["Risk", "Medium", "Low"]],
         "bottom_takeaway": "So what."},
        {"slide_number": 11, "layout_type": "closing", "action_title": "Thank You", "subtitle_or_context": "Presented by Test"},
    ]


def check_builtin_decks() -> None:
    from chatppt.pptx_build import build_presentation_pptx
    from chatppt.presets import BUILTIN_THEMES

    slides = sample_slides()
    # Two themes are enough to catch theme-specific crashes; re-opening the bytes proves the file is valid
    for theme in list(BUILTIN_THEMES)[:2]:
        data = build_presentation_pptx(slides, theme, None, False, "", [])
        out = Presentation(io.BytesIO(data))
        assert len(out.slides) == len(slides), theme
    # comparison_table renders a NATIVE table with the header row and every cell
    table_slide = out.slides[9]
    tables = [sh.table for sh in table_slide.shapes if sh.has_table]
    assert tables and len(tables[0].rows) == 4 and len(tables[0].columns) == 3
    assert tables[0].cell(0, 1).text == "Option A" and tables[0].cell(3, 2).text == "Low"
    # timeline_flow renders every stage label
    texts = {sh.text_frame.text for sh in out.slides[8].shapes if sh.has_text_frame}
    assert {"Q1", "Q4", "1", "4", "Step 1\nWhat happens."} <= texts, texts
    print(f"[ok] built-in decks: {len(slides)} slides, all 11 layouts (native table, timeline nodes)")


def _make_template() -> bytes:
    """Build a 16:9 user template that has a border frame on its master and explicit fonts on a sample slide."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[6]  # "Blank" layout
    tmp = prs.slides.add_slide(layout)  # the template's one sample slide (explicit Georgia / Verdana runs)
    frame = tmp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(12.3), Inches(6.5))
    frame.fill.background()
    frame.name = "TemplateFrame"
    # python-pptx can't add shapes to a master directly, so move the frame's XML from the slide onto the master.
    # It then shows on every slide, like the decorative borders in real corporate templates.
    el = frame._element
    el.getparent().remove(el)
    prs.slide_masters[0].shapes._spTree.append(el)
    # A large title run and a small body run; the builder should pick these up as the template's title/body fonts
    for name, size, text in (("Georgia", 36, "Big title"), ("Verdana", 14, "body text")):
        tb = tmp.shapes.add_textbox(Inches(1), Inches(1 if size > 20 else 3), Inches(8), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text, run.font.size, run.font.name = text, Pt(size), name
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def check_custom_template() -> None:
    from chatppt.pptx_build import build_presentation_pptx

    data = build_presentation_pptx(
        sample_slides()[:3], "Corporate Blue", None, False, "", [], custom_template_bytes=_make_template()
    )
    out = Presentation(io.BytesIO(data))
    # Only placeholders may remain on the master; decorations like the frame would overlap the generated layouts
    leftovers = [s.name for s in out.slide_masters[0].shapes if not s.is_placeholder]
    assert not leftovers, f"master decorations still present: {leftovers}"
    assert len(out.slides) == 3
    fonts = {
        run.font.name
        for slide in out.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
    }
    assert {"Georgia", "Verdana"} <= fonts, f"template fonts not applied: {fonts}"
    print("[ok] custom template: border frame stripped, template fonts applied")


def check_helpers() -> None:
    from chatppt.llm_client import extract_clean_json_str
    from chatppt.outline import mentions_slide_count, parse_slide_count_change
    from chatppt.ui.editor_text import format_points_for_editor, parse_points_from_editor

    pts = [{"heading": "A", "detail": "B"}, {"heading": "C", "detail": "D"}]
    assert parse_points_from_editor(format_points_for_editor(pts)) == pts
    assert extract_clean_json_str('<think>hmm</think>```json\n{"a": 1}\n```') == '{"a": 1}'
    assert mentions_slide_count("make it 6 slides") and not mentions_slide_count("shorter titles")
    # Relative slide-count edits: new target = current deck size + delta (regression: "add one more slide" used to be trimmed back)
    assert parse_slide_count_change("add one more slide conclusion", 8) == 9
    assert parse_slide_count_change("add 2 more slides", 8) == 10
    assert parse_slide_count_change("add a conclusion slide", 8) == 9
    assert parse_slide_count_change("remove the last slide", 8) == 7
    assert parse_slide_count_change("\u52a0\u4e00\u9875\u7ed3\u8bba", 8) == 9
    assert parse_slide_count_change("make it 6 slides", 8) is None
    assert parse_slide_count_change("add a chart to slide 3", 8) is None
    assert parse_slide_count_change("shorter titles", 8) is None
    assert parse_slide_count_change("add a thank you slide", 8) == 9
    assert parse_slide_count_change("add thank you slide", 8) == 9
    assert parse_slide_count_change("remove the thank you slide", 8) == 7
    assert parse_slide_count_change("remove duplicates from slide 3", 8) is None
    print("[ok] helpers")


def check_closing_slide() -> None:
    from chatppt.outline import _trim_slides_to_count, sanitize_outline_dict
    from chatppt.schemas import LayoutType, PresentationOutline, SlideSpec

    def spec(i, layout):
        return SlideSpec(slide_number=i + 1, layout_type=layout, action_title=f"T{i}")

    # A misplaced closing slide is moved to the end, kept unique and given no takeaway banner
    layouts = ["cover", "closing", "bullet_cards", "bullet_cards", "bullet_cards"]
    out = sanitize_outline_dict(PresentationOutline(deck_title="d", slides=[spec(i, lt) for i, lt in enumerate(layouts)]))
    kinds = [s["layout_type"] for s in out]
    assert kinds[-1] == "closing" and kinds.count("closing") == 1, kinds
    assert out[-1]["bottom_takeaway"] is None
    # Trimming never drops the closing slide or the conclusion right before it
    deck = [spec(i, lt) for i, lt in enumerate(["cover", "bullet_cards", "bullet_cards", "bullet_cards", "closing"])]
    trimmed = _trim_slides_to_count(deck, 4)
    assert len(trimmed) == 4 and trimmed[-1].layout_type == LayoutType.CLOSING and trimmed[-2].action_title == "T3"
    print("[ok] closing slide: placement, no takeaway, trim-safe")


def check_new_layouts() -> None:
    from chatppt.outline import normalize_table, sanitize_outline_dict
    from chatppt.schemas import PresentationOutline, SlideSpec
    from chatppt.ui.editor_text import format_steps_for_editor, format_table_for_editor, parse_steps_from_editor, parse_table_from_editor

    # Ragged rows are padded / cut to the header count; a 1-column table is rejected
    headers, rows = normalize_table(["A", "B", "C"], [["1"], ["1", "2", "3", "4"], ["", ""]])
    assert headers == ["A", "B", "C"] and rows == [["1", "", ""], ["1", "2", "3"]]
    assert normalize_table(["Only"], [["x"]]) == ([], [])
    pts = [{"heading": "H1", "detail": "D1"}, {"heading": "H2", "detail": "D2"}, {"heading": "H3", "detail": "D3"}]
    slides = [
        SlideSpec(slide_number=1, layout_type="cover", action_title="c"),
        SlideSpec(slide_number=2, layout_type="comparison_table", action_title="bad table", points=pts),
        SlideSpec(slide_number=3, layout_type="timeline_flow", action_title="steps from points", points=pts),
        SlideSpec(slide_number=4, layout_type="timeline_flow", action_title="empty timeline"),
    ]
    out = sanitize_outline_dict(PresentationOutline(deck_title="d", slides=slides))
    assert out[1]["layout_type"] != "comparison_table"  # no table data -> falls back to cards
    assert out[2]["layout_type"] == "timeline_flow" and [s["title"] for s in out[2]["timeline_steps"]] == ["H1", "H2", "H3"]
    assert out[3]["layout_type"] != "timeline_flow"
    steps = [{"stage_label": "Q1", "title": "Pilot", "description": "Run it"}]
    assert parse_steps_from_editor(format_steps_for_editor(steps)) == steps
    assert parse_table_from_editor(format_table_for_editor(["A", "B"], [["x", "y"]])) == (["A", "B"], [["x", "y"]])
    print("[ok] timeline / table: sanitizer fallbacks, rectangular tables, editor round-trip")


def check_report_images() -> None:
    from PIL import Image

    from chatppt.extraction import prepare_report_image
    from chatppt.images import assign_report_images, resolve_slide_image_bytes
    from chatppt.schemas import IMAGE_MODE_PLACEHOLDER, IMAGE_MODE_REPORT_FIRST

    def png(w, h, mode="RGB"):
        buf = io.BytesIO()
        Image.new(mode, (w, h), "white").save(buf, format="PNG")
        return buf.getvalue()

    assert prepare_report_image(png(64, 64)) is None  # logo / icon size is skipped
    assert prepare_report_image(b"not an image") is None
    big = prepare_report_image(png(3000, 1500))
    assert big and max(big["width"], big["height"]) <= 2400  # huge scans are downscaled
    assert prepare_report_image(png(400, 300, "RGBA"))["bytes"][1:4] == b"PNG"  # alpha kept as PNG

    rep, img, txt = IMAGE_MODE_REPORT_FIRST, "image_split", "bullet_cards"
    deck = [{"layout_type": txt}, {"layout_type": img, "image_source_mode": rep},
            {"layout_type": img, "image_source_mode": rep, "report_image_index": 0},
            {"layout_type": img, "image_source_mode": IMAGE_MODE_PLACEHOLDER},
            {"layout_type": img, "image_source_mode": rep}]
    # Explicit pick wins, the others get the remaining figures in order; placeholder slides are skipped
    assert assign_report_images(deck, 3, rep) == {2: 0, 1: 1, 4: 2}
    assert assign_report_images(deck, 1, rep) == {2: 0}  # out of figures -> others fall back
    images = [{"name": "a", "bytes": b"A"}, {"name": "b", "bytes": b"B"}]
    slide = {"layout_type": img, "image_source_mode": rep, "report_image_index": 1}
    assert resolve_slide_image_bytes(slide, None, False, "", images) == b"B"
    print("[ok] report images: size filter, downscale, assignment, resolver")


def check_assistant() -> None:
    from chatppt.chat_assistant import build_assistant_prompt, parse_assistant_reply

    assert parse_assistant_reply("ACTION: REVISE").action == "revise"
    assert parse_assistant_reply("<think>hmm</think>ACTION: REVISE\n").action == "revise"
    dec = parse_assistant_reply("ACTION: ANSWER\nBecause the data is a comparison.")
    assert dec.action == "answer" and dec.reply == "Because the data is a comparison."
    assert parse_assistant_reply("**ACTION: ANSWER**\nHi there").reply == "Hi there"
    assert parse_assistant_reply("Just a plain answer").action == "answer"
    assert parse_assistant_reply("").action == "revise" and parse_assistant_reply("ACTION: ANSWER").action == "revise"
    prompt = build_assistant_prompt(
        "why three cards?", "SOURCE_MARKER text", sample_slides(), [{"role": "user", "content": "HISTORY_MARKER"}]
    )
    for marker in ("SOURCE_MARKER", "Three cards title", "KPI title", "HISTORY_MARKER", "why three cards?"):
        assert marker in prompt, marker
    # Targeted edits: PATCH with valid slide numbers; anything unusable degrades to a full revise
    dec = parse_assistant_reply("ACTION: PATCH: 3, 5", deck_size=9)
    assert dec.action == "patch" and dec.slides == [3, 5], dec
    assert parse_assistant_reply("**ACTION: PATCH 2-4**", deck_size=9).slides == [2, 3, 4]
    assert parse_assistant_reply("ACTION: PATCH: 12", deck_size=9).action == "revise"
    assert parse_assistant_reply("ACTION: PATCH:", deck_size=9).action == "revise"
    assert parse_assistant_reply("ACTION: PATCH: 1-9", deck_size=9).action == "revise"  # too many slides
    # A long source is cut to relevant excerpts: a fact far past the first 12k chars must still reach the model
    long_source = "Intro.\n\n" + "\n\n".join("Filler paragraph about logistics. " * 20 for _ in range(60))
    long_source += "\n\nChurn fell to 3.2% after the NEEDLE loyalty program.\n\n" + "Tail filler. " * 400
    prompt = build_assistant_prompt("what happened to churn after the loyalty program?", long_source, [], [])
    assert "NEEDLE" in prompt and len(prompt) < 14000
    print("[ok] assistant routing + prompt context + patch protocol")


def check_retrieval() -> None:
    from chatppt.retrieval import select_relevant_context, split_into_chunks, tokenize

    assert "增长" in tokenize("营收增长62%") and "62%" in tokenize("营收增长62%")
    assert all(len(c) <= 1000 for c in split_into_chunks("x" * 5000))
    assert select_relevant_context("short text", "anything", 100) == "short text"
    doc = "\n\n".join(f"Section {i}: generic operations text. " * 12 for i in range(40))
    doc += "\n\n最终客户流失率下降到 3.2%。"
    out = select_relevant_context(doc, "客户流失率是多少", 3000)
    assert "3.2%" in out and len(out) <= 3000 + 20 and out.startswith("Section 0")
    # Nothing matches -> head of the document, like the old truncation
    assert select_relevant_context(doc, "zzz", 3000).startswith(doc[:2000])
    print("[ok] retrieval: CJK bigrams, chunking, BM25 recall, head fallback")


def check_patch() -> None:
    import chatppt.outline as outline
    from chatppt.schemas import SlidePatch, SlideSpec

    deck = sample_slides()
    deck[3]["image_source_mode"] = "Reserve Photo Placeholder"
    new = SlideSpec(slide_number=4, layout_type="bullet_cards", action_title="Patched title",
                    points=[{"heading": "H", "detail": "D"}])
    stray_closing = SlideSpec(slide_number=5, layout_type="closing", action_title="Bye")
    # Stub out the LLM call so the patch logic runs offline; the "model" returns slide 4 rewritten plus a
    # misplaced closing slide for slide 5. Always restore the real function, even if the patch raises.
    real = outline._request_outline
    try:
        outline._request_outline = lambda *a, **k: SlidePatch(slides=[new, stray_closing])
        out = outline.patch_slides("Groq", "k", "m", None, deck, [4, 5], "rewrite", "source", "tone", "English")
    finally:
        outline._request_outline = real
    assert len(out) == len(deck)
    # Untouched slides are the very same objects: a patch can never drift them
    assert all(out[i] is deck[i] for i in range(len(deck)) if i not in (3, 4))
    assert out[3]["action_title"] == "Patched title" and out[3]["slide_number"] == 4
    assert out[3]["image_source_mode"] == "Reserve Photo Placeholder"  # user's per-slide choice kept
    assert out[4]["layout_type"] == "bullet_cards"  # closing is only allowed as the last slide
    print("[ok] targeted patch: only target slides replaced, closing guard, image mode kept")


def check_history_and_projects() -> None:
    from chatppt.deck_ops import HISTORY_MAX, can_redo, can_undo, push_history, redo, undo
    from chatppt.project_io import ProjectFileError, export_project, load_project

    chat = {"slides": sample_slides()}
    push_history(chat)
    chat["slides"][2]["action_title"] = "Edited after snapshot"  # snapshots must be deep copies
    v2 = chat["slides"]
    assert can_undo(chat) and not can_redo(chat)
    assert undo(chat)[2]["action_title"] == "Bullet cards title" and can_redo(chat)
    assert redo(chat) is not None and chat["slides"][2]["action_title"] == "Edited after snapshot" and chat["slides"] == v2
    push_history(chat)
    undo(chat)
    push_history(chat)  # a new change clears the redo stack
    assert not can_redo(chat)
    for _ in range(HISTORY_MAX + 5):
        push_history(chat)
    assert len(chat["slides_history"]) == HISTORY_MAX
    assert undo({"slides": None}) is None

    # Project round trip: outline, messages, figures and safe settings survive; secrets never leave
    deck = sample_slides()
    deck[7]["image_source_mode"], deck[7]["report_image_index"] = "Report Images First (Fallback to Unsplash / AI)", 0
    src_chat = {"title": "Q3 review", "slides": deck, "source_text": "SRC", "messages": [{"role": "user", "content": "hi"}],
                "report_images": [{"name": "fig", "width": 400, "height": 300, "bytes": b"\x89PNGdata"}]}
    settings = {"api_keys": {"Groq": "gsk_SECRET"}, "unsplash_key": "UNSPLASH_SECRET", "provider": "Google Gemini",
                "models": {"Google Gemini": "gemini-x"}, "image_mode": "Reserve Photo Placeholder", "use_imagen": True}
    raw = export_project(src_chat, settings, {"opt_slides": 11, "opt_theme": "Corporate Blue", "opt_tone": "nope"})
    assert b"SECRET" not in raw
    loaded = load_project(raw)
    out = loaded["chat"]
    assert [s["action_title"] for s in out["slides"]] == [s["action_title"] for s in deck]
    assert out["slides"][9]["table_rows"] == deck[9]["table_rows"] and out["slides"][7]["report_image_index"] == 0
    assert out["report_images"][0]["bytes"] == b"\x89PNGdata" and out["source_text"] == "SRC"
    assert loaded["settings"] == {"provider": "Google Gemini", "models": {"Google Gemini": "gemini-x"},
                                  "image_mode": "Reserve Photo Placeholder", "use_imagen": True}
    assert loaded["options"] == {"opt_slides": 11, "opt_theme": "Corporate Blue"}  # invalid tone dropped
    for bad in (b"not json", b'{"format": "other"}', raw.replace(b'"layout_type": "cover"', b'"layout_type": "bogus"', 1)):
        try:
            load_project(bad)
            raise AssertionError("expected ProjectFileError")
        except ProjectFileError:
            pass
    print("[ok] history (deep snapshots, redo cleared, capped) + project save/load (no secrets, validated)")


def check_structure_ops() -> None:
    from chatppt.deck_ops import can_delete, can_move, delete_slide, insert_slide, move_slide

    deck = sample_slides()  # 11 slides: cover ... closing
    titles = lambda d: [s["action_title"] for s in d]  # noqa: E731
    moved = move_slide(deck, 3, -1)
    assert titles(moved)[2:4] == ["Three cards title", "Bullet cards title"]
    assert [s["slide_number"] for s in moved] == list(range(1, 12))
    # The cover stays first and the closing slide stays last
    assert not can_move(deck, 1, -1) and not can_move(deck, 0, 1) and not can_move(deck, 9, 1) and not can_move(deck, 10, -1)
    assert move_slide(deck, 1, -1)[0]["layout_type"] == "cover"
    assert not can_delete(deck, 0) and can_delete(deck, 10)
    shorter = delete_slide(deck, 2)
    assert len(shorter) == 10 and "Bullet cards title" not in titles(shorter) and shorter[-1]["slide_number"] == 10
    # Inserting after the closing slide still lands before it; the blank slide is fully sanitized
    grown, pos = insert_slide(deck, 10, "Reserve Photo Placeholder")
    assert len(grown) == 12 and pos == 10 and grown[-1]["layout_type"] == "closing"
    assert grown[pos]["layout_type"] == "bullet_cards" and grown[pos]["kpis"] and grown[pos]["table_rows"]
    assert len(insert_slide([dict(s) for s in deck * 2][:20], 3, "x")[0]) == 20  # 20-slide cap
    print("[ok] structure ops: move/insert/delete keep cover first and closing last, renumbered, capped")


def check_preview() -> None:
    import copy

    from chatppt.preview import overflows, render_slide_preview
    from chatppt.presets import BUILTIN_THEMES

    theme = BUILTIN_THEMES["Corporate Blue"]
    deck = copy.deepcopy(sample_slides())
    for i, s in enumerate(deck):
        html_out, warnings = render_slide_preview(s, theme, i, len(deck))
        assert html_out.startswith('<div class="cpv-slide"') and not warnings, (s["layout_type"], warnings)
    # Overflow is caught for long text (Latin and CJK) and short text is not flagged
    assert not overflows([{"text": "Short heading", "pt": 17}], 10.5, 1.2)
    long_detail = copy.deepcopy(deck[2])
    long_detail["points"][0]["detail"] = "很长的说明文字。" * 30
    assert render_slide_preview(long_detail, theme, 2, 11)[1] == ["Card 1 text overflows its card."]
    long_title = copy.deepcopy(deck[3])
    long_title["action_title"] = "A very long action title that keeps going " * 4
    assert "Title is too long" in render_slide_preview(long_title, theme, 3, 11)[1][0]
    # User text is escaped (st.html is not iframed)
    evil = dict(deck[2], action_title="<script>alert(1)</script>")
    assert "<script>" not in render_slide_preview(evil, theme, 2, 11)[0]
    print("[ok] preview: all 11 layouts render, overflow estimate (Latin + CJK), HTML escaped")


def check_browser_prefs() -> None:
    from chatppt.prefs import collect_prefs, parse_prefs, prefs_json

    settings = {"provider": "Google Gemini", "models": {"Google Gemini": "gemini-x"}, "image_mode": "Reserve Photo Placeholder",
                "use_imagen": True, "api_keys": {"Google Gemini": "AIza-KEY", "Groq (Fast & Free Tier)": ""},
                "unsplash_key": "UNSPLASH-KEY", "remember_keys": True}
    options = {"opt_slides": 12, "opt_theme": "Modern Dark", "opt_language": "English", "opt_tone": None}
    raw = prefs_json(collect_prefs(settings, options))
    back = parse_prefs(raw)
    assert back["api_keys"] == {"Google Gemini": "AIza-KEY"} and back["unsplash_key"] == "UNSPLASH-KEY"
    assert back["settings"]["provider"] == "Google Gemini" and back["options"] == {"opt_slides": 12, "opt_theme": "Modern Dark",
                                                                                   "opt_language": "English"}
    assert prefs_json(collect_prefs(settings, options)) == raw  # stable: unchanged prefs are never rewritten
    # Remember-keys off: no secret of any kind is in the payload
    no_keys = prefs_json(collect_prefs(dict(settings, remember_keys=False), options))
    assert "KEY" not in no_keys and parse_prefs(no_keys)["remember_keys"] is False
    # Garbage, other versions, unknown providers and oversized values are dropped
    for bad in (None, "", "not json", "[]", '{"v": 99}'):
        assert parse_prefs(bad) is None
    hostile = parse_prefs('{"v": 1, "api_keys": {"Evil": "x", "Google Gemini": "' + "A" * 600 + '"}, '
                          '"settings": {"provider": "Evil", "image_mode": "x"}, "options": {"opt_slides": 999}}')
    assert hostile["api_keys"] == {} and hostile["settings"] == {} and hostile["options"] == {"opt_slides": 20}
    print("[ok] browser prefs: round trip, stable payload, no keys when off, hostile payload sanitized")


def check_chat_store() -> None:
    import copy

    import chatppt.chat_store as cs

    chat = {"id": "abc123", "title": "Q3 review", "messages": [{"role": "user", "content": "make a deck"}],
            "slides": copy.deepcopy(sample_slides()), "source_text": "SOURCE " * 100, "pending": None, "updated_at": 5.0,
            "report_images": [{"name": "fig", "width": 400, "height": 300, "bytes": b"\x89PNGdata"}]}
    records = {cs.record_key(chat["id"], p): cs.part_payload(chat, p) for p in cs.PARTS}
    back = cs.restore_chats(records)[0]
    assert back["id"] == "abc123" and back["title"] == "Q3 review" and back["source_text"] == chat["source_text"]
    assert [s["action_title"] for s in back["slides"]] == [s["action_title"] for s in chat["slides"]]
    assert back["report_images"][0]["bytes"] == b"\x89PNGdata" and back["updated_at"] == 5.0
    # Only the part that changed gets a new fingerprint (a slide edit never re-sends source or figures)
    before = {p: cs.part_signature(chat, p) for p in cs.PARTS}
    chat["slides"][2]["action_title"] = "Edited"
    after = {p: cs.part_signature(chat, p) for p in cs.PARTS}
    assert after["chat"] != before["chat"] and after["source"] == before["source"] and after["images"] == before["images"]
    # Oversized figures are not kept in the browser (the chat still is)
    big = dict(chat, report_images=[{"name": "huge", "bytes": b"x" * (cs.MAX_IMAGES_BYTES + 1)}])
    assert '"report_images":[]' in cs.part_payload(big, "images")
    # Broken records are skipped; newest first; capped
    records["zzz/chat"] = "not json"
    records["bad/chat"] = '{"v": 1, "messages": [], "slides": [{"layout_type": "bogus"}]}'
    many = dict(records)
    for i in range(cs.MAX_STORED_CHATS + 3):
        many[f"c{i}/chat"] = cs.part_payload({"title": f"T{i}", "messages": [{"role": "user", "content": "x"}],
                                              "slides": [], "updated_at": float(i)}, "chat")
    restored = cs.restore_chats(many)
    assert len(restored) == cs.MAX_STORED_CHATS and restored[0]["title"] == f"T{cs.MAX_STORED_CHATS + 2}"
    assert all(c["id"] not in ("zzz", "bad") for c in restored)
    assert not cs.is_storable({"messages": [], "slides": None, "pending": None})
    assert not cs.is_storable({"messages": [{"role": "user", "content": "x"}], "pending": {"job": 1}})
    print("[ok] chat store: 3-part round trip, per-part change detection, size cap, broken records skipped")


def check_outline_anchor() -> None:
    from chatppt.project_io import export_project, load_project
    from chatppt.ui.messages import deck_anchor, message_kind

    deck_msg = {"role": "assistant", "content": "Your outline is ready — 9 slides.", "kind": "deck"}
    answer = {"role": "assistant", "content": "Because it compares two options.", "kind": "answer"}
    error = {"role": "assistant", "content": ":material/error: I couldn't generate the outline.", "kind": "error"}
    user = {"role": "user", "content": "why?"}
    # Questions and failed requests after the deck never move the outline
    assert deck_anchor([user, deck_msg, user, answer, user, error]) == 1
    patched = {"role": "assistant", "content": "Updated **slide 3**; all other slides are unchanged.", "kind": "deck"}
    assert deck_anchor([user, deck_msg, user, answer, user, patched, user, answer]) == 5
    assert deck_anchor([user, answer]) is None  # no deck message (e.g. an old project) -> outline at the end
    # Messages saved before "kind" existed are recognized by their text
    assert message_kind({"role": "assistant", "content": "Your outline is ready — **8 slides**."}) == "deck"
    assert message_kind({"role": "assistant", "content": "Updated **slides 2, 5**; all other slides..."}) == "deck"
    assert message_kind({"role": "assistant", "content": "It means...", "meta": "💬 Answer · outline unchanged"}) == "answer"
    assert message_kind({"role": "assistant", "content": "❌ I couldn't generate the outline."}) == "error"
    # The kind survives a project round trip
    raw = export_project({"title": "t", "slides": sample_slides(), "messages": [user, deck_msg, user, answer]}, {}, {})
    assert [m.get("kind") for m in load_project(raw)["chat"]["messages"]] == [None, "deck", None, "answer"]
    print("[ok] outline anchor: stays after the last deck change; questions/errors don't move it; legacy messages")


def check_token_limits() -> None:
    import json
    import urllib.error
    import urllib.request

    import chatppt.llm_client as llm

    # Builds an HTTPError whose JSON body copies Groq's real 413/429 token-limit message, which the client parses
    def groq_error(code, limit, requested, used=None, retry=None):
        used_part = f"Used {used}, " if used is not None else ""
        retry_part = f" Please try again in {retry}s." if retry is not None else ""
        body = json.dumps({"error": {"message": f"Request too large for model `m` on tokens per minute (TPM): "
                                                f"Limit {limit}, {used_part}Requested {requested}.{retry_part}",
                                     "type": "tokens", "code": "rate_limit_exceeded"}})
        return urllib.error.HTTPError("u", code, "err", {}, io.BytesIO(body.encode()))

    # A minimal successful urlopen() response (used as a context manager, read() returns a chat completion)
    class Ok:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self):
            return json.dumps({"choices": [{"message": {"content": '{"ok": 1}'}, "finish_reason": "stop"}]}).encode()

    # Sends one request through the real client with urlopen replaced: each attempt pops the next scripted
    # response and records the max_completion_tokens it asked for, so retries and cap shrinking can be checked
    def run(responses):
        caps, real = [], urllib.request.urlopen

        def fake(req, timeout=None):
            caps.append(json.loads(req.data).get("max_completion_tokens"))
            r = responses.pop(0)
            if isinstance(r, Exception):
                raise r
            return r

        urllib.request.urlopen = fake
        try:
            return llm._post_chat_completion("https://api.groq.com/x", {}, {"max_completion_tokens": 8192}, "m"), caps
        finally:
            urllib.request.urlopen = real

    # The user's case: input alone (~8.4k) exceeds an 8k TPM plan -> one readable error, no retries
    try:
        run([groq_error(413, 8000, 16550)])
        raise AssertionError("expected TokenBudgetError")
    except llm.TokenBudgetError as e:
        assert "too large" in str(e) and "8,000" in str(e)
    # Only the reserved output makes it too big -> retried once with a smaller output cap
    _, caps = run([groq_error(413, 12000, 15300), Ok()])
    # New cap = limit - input tokens (requested - old cap) - 100 safety margin
    assert caps[0] == 8192 and caps[1] == 12000 - (15300 - 8192) - 100, caps
    # Minute budget temporarily used up -> short wait, then the same request again
    _, caps = run([groq_error(429, 12000, 4000, used=9000, retry=0.01), Ok()])
    assert caps == [8192, 8192]
    print("[ok] token limits: oversized request explained, output cap shrink, short rate-limit wait")


if __name__ == "__main__":
    # Each check prints "[ok] ..." when it passes; the first failed assert stops the run with a traceback and exit code 1
    for fn in (check_imports, check_builtin_decks, check_custom_template, check_helpers, check_closing_slide,
               check_new_layouts, check_report_images, check_assistant, check_retrieval, check_patch, check_token_limits,
               check_history_and_projects, check_structure_ops, check_preview,
               check_browser_prefs, check_chat_store, check_outline_anchor):
        fn()
    print("ALL SMOKE TESTS PASSED")
    sys.exit(0)